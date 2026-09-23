"""
Act 1 artifact — "this is what you do today".

Builds a synthetic, hand-built-looking Excel rate-indication workbook for
General Liability / Germany 2027 that lands on the SAME +6.6% the Databricks
notebook, tables and app reproduce later. Live formulas (not pasted values) so
it reads like the real fragile spreadsheet an actuary maintains — and so you can
wiggle an assumption and watch the number move, just like the app does later.

Run: python3 scripts/build_act1_workbook.py  ->  data/Indication_GL_DE_2027_v7_FINAL.xlsx
"""
import os, sys
sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "src", "app"))
sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "src", "common"))
from indication_engine import ExperienceYear, calc_segment  # noqa
from gen_book import generate, baseline_assumptions, _ldf_for_maturity  # noqa

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment

LOB, TERR, PERIOD, TAIL = "GENERAL_LIABILITY", "DE", 2027, "long"
book = generate("eu_commercial")
rows = sorted([e for e in book["experience"] if e["lob_code"] == LOB and e["territory_code"] == TERR],
              key=lambda e: e["accident_year"])
crl = {(r["lob_code"], r["territory_code"]): r for r in book["rate_state"]}[(LOB, TERR)]["current_rate_level"]
A = baseline_assumptions(TAIL, 0.055, -0.015)          # the approved baseline basis
n = int(A["experience_period_years"])
exp = rows[-n:]                                         # last 5 accident years
latest_base_ldf = _ldf_for_maturity(TAIL, 0)           # 1.22

# sanity: what the platform will reproduce
ey = [ExperienceYear(accident_year=r["accident_year"], earned_premium=r["earned_premium"],
                     reported_incurred=r["reported_incurred"], claim_count=r["claim_count"],
                     exposure=r["exposure"], rate_level_index=r["rate_level_index"],
                     ldf_to_ultimate=r["ldf_to_ultimate"]) for r in rows]
target = calc_segment(ey, crl, PERIOD, A).indicated_rate_change
print(f"engine indicated (target the sheet must match): {target*100:+.2f}%")

wb = Workbook()
HDR = Font(bold=True, color="FFFFFF"); HDRFILL = PatternFill("solid", fgColor="1F3864")
YEL = PatternFill("solid", fgColor="FFF2CC"); BOLD = Font(bold=True)
MONEY = "#,##0"; PCT = "0.0%"; FAC = "0.000"
thin = Border(*(Side(style="thin", color="D9D9D9"),) * 4)
right = Alignment(horizontal="right")

# ---- Assumptions tab (the hand-keyed knobs) ----
a = wb.active; a.title = "Assumptions"
a["A1"] = "GL / Germany — 2027 rate indication — ASSUMPTIONS"; a["A1"].font = Font(bold=True, size=13)
a["A2"] = "Bricksurance SE (SYNTHETIC / illustrative)"; a["A2"].font = Font(italic=True, color="808080")
akeys = [
    ("Severity trend", A["severity_trend"], PCT, "gut feel off latest large claims — check with claims?"),
    ("Frequency trend", A["frequency_trend"], PCT, ""),
    ("Selected development factor (tail)", A["loss_development_factor"], FAC, "picked off the triangle tab, roughly"),
    ("Large-loss load", A["large_loss_load"], PCT, ""),
    ("Catastrophe load", A["cat_load"], PCT, ""),
    ("Credibility (Z)", A["credibility"], PCT, "we always use 0.75 for GL"),
    ("Experience period (years)", A["experience_period_years"], "0", ""),
    ("Expense ratio", A["expense_ratio"], PCT, ""),
    ("Commission ratio", A["commission_ratio"], PCT, ""),
    ("Reinsurance load", A["reinsurance_load"], PCT, ""),
    ("Profit & contingency", A["profit_provision"], PCT, ""),
    ("Current rate level index", crl, FAC, ""),
    ("Latest-year empirical LDF", latest_base_ldf, FAC, ""),
    ("Prospective period", PERIOD, "0", ""),
]
for i, (label, val, fmt, note) in enumerate(akeys, start=4):
    a[f"A{i}"] = label; a[f"B{i}"] = val; a[f"B{i}"].number_format = fmt; a[f"B{i}"].fill = YEL
    if note:
        a[f"B{i}"].comment = Comment(note, "actuary")
a.column_dimensions["A"].width = 34; a.column_dimensions["B"].width = 14
# named refs (row numbers)
R = {label: i for i, (label, *_ ) in enumerate(akeys, start=4)}
sev, freq, seldf = f"Assumptions!$B${R['Severity trend']}", f"Assumptions!$B${R['Frequency trend']}", f"Assumptions!$B${R['Selected development factor (tail)']}"
ll, cat, cred = f"Assumptions!$B${R['Large-loss load']}", f"Assumptions!$B${R['Catastrophe load']}", f"Assumptions!$B${R['Credibility (Z)']}"
exp_r, comm, ri = f"Assumptions!$B${R['Expense ratio']}", f"Assumptions!$B${R['Commission ratio']}", f"Assumptions!$B${R['Reinsurance load']}"
profit, crlref, latest, prosp = f"Assumptions!$B${R['Profit & contingency']}", f"Assumptions!$B${R['Current rate level index']}", f"Assumptions!$B${R['Latest-year empirical LDF']}", f"Assumptions!$B${R['Prospective period']}"

# ---- Experience tab (inputs + the fragile formula chain) ----
e = wb.create_sheet("Experience")
cols = ["Accident year", "Earned premium", "Reported incurred", "Rate index", "Empirical LDF",
        "On-level premium", "Eff. LDF", "Ultimate", "Trend yrs", "Trend factor", "Trended ultimate", "Loss ratio"]
for j, c in enumerate(cols, start=1):
    cell = e.cell(row=1, column=j, value=c); cell.font = HDR; cell.fill = HDRFILL; cell.alignment = right
r0 = 2
for k, row in enumerate(exp):
    r = r0 + k; ay = row["accident_year"]; base_ldf = _ldf_for_maturity(TAIL, 2025 - ay)
    e.cell(row=r, column=1, value=ay)
    e.cell(row=r, column=2, value=round(row["earned_premium"], 0)).number_format = MONEY
    e.cell(row=r, column=3, value=round(row["reported_incurred"], 0)).number_format = MONEY
    e.cell(row=r, column=4, value=round(row["rate_level_index"], 4)).number_format = FAC
    e.cell(row=r, column=5, value=round(base_ldf, 4)).number_format = FAC
    e.cell(row=r, column=6, value=f"=B{r}*({crlref}/D{r})").number_format = MONEY          # on-level
    e.cell(row=r, column=7, value=f"=1+(E{r}-1)*({seldf}/{latest})").number_format = FAC    # eff ldf
    e.cell(row=r, column=8, value=f"=C{r}*G{r}").number_format = MONEY                       # ultimate
    e.cell(row=r, column=9, value=f"={prosp}-A{r}")                                          # trend yrs
    e.cell(row=r, column=10, value=f"=(1+{freq})^I{r}*(1+{sev})^I{r}").number_format = FAC    # trend factor
    e.cell(row=r, column=11, value=f"=H{r}*J{r}").number_format = MONEY                       # trended
    e.cell(row=r, column=12, value=f"=K{r}/F{r}").number_format = PCT                         # LR
rn = r0 + len(exp)
e.cell(row=rn, column=1, value="Total").font = BOLD
e.cell(row=rn, column=6, value=f"=SUM(F{r0}:F{rn-1})").number_format = MONEY; e.cell(row=rn, column=6).font = BOLD
e.cell(row=rn, column=11, value=f"=SUM(K{r0}:K{rn-1})").number_format = MONEY; e.cell(row=rn, column=11).font = BOLD
for col, w in zip("ABCDEFGHIJKL", [13, 15, 16, 10, 12, 15, 9, 14, 9, 12, 15, 10]):
    e.column_dimensions[col].width = w
SUMF, SUMK = f"Experience!$F${rn}", f"Experience!$K${rn}"

# ---- Indication tab (the number that gets emailed round) ----
ind = wb.create_sheet("Indication", 0)
ind["A1"] = "GL / GERMANY — 2027 RATE INDICATION"; ind["A1"].font = Font(bold=True, size=14)
ind["A2"] = "prepared by: a.actuary   ·   file: Indication_GL_DE_2027_v7_FINAL.xlsx   ·   SYNTHETIC"; ind["A2"].font = Font(italic=True, color="808080")
steps = [
    ("Experience loss ratio (on-level, developed, trended)", f"={SUMK}/{SUMF}", PCT),
    ("+ loads (large-loss, cat)", f"=B4*(1+{ll})+{cat}", PCT),
    ("Permissible loss ratio  = 1 - (expenses+commission+RI+profit)", f"=1-({exp_r}+{comm}+{ri}+{profit})", PCT),
    ("Credibility-weighted projected loss ratio", f"={cred}*B5+(1-{cred})*B6", PCT),
    ("INDICATED RATE CHANGE  = projected / permissible - 1", "=B7/B6-1", PCT),
]
for i, (label, formula, fmt) in enumerate(steps, start=4):
    ind[f"A{i}"] = label; ind[f"B{i}"] = formula; ind[f"B{i}"].number_format = fmt
ind["A8"].font = BOLD; ind["B8"].font = Font(bold=True, size=13); ind["B8"].fill = YEL
ind["A10"] = "Selected rate change (mgmt): "; ind["B10"] = 0.0; ind["B10"].number_format = PCT; ind["B10"].fill = YEL
ind["A11"] = "Sign-off:"; ind["B11"] = "(pending — see email thread)"; ind["B11"].font = Font(italic=True, color="C00000")
ind.column_dimensions["A"].width = 52; ind.column_dimensions["B"].width = 16

os.makedirs(os.path.join(os.path.dirname(__file__), "..", "data"), exist_ok=True)
out = os.path.join(os.path.dirname(__file__), "..", "data", "Indication_GL_DE_2027_v7_FINAL.xlsx")
wb.save(out)
print("wrote", out)
