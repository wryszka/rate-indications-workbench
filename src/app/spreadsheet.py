"""
Spreadsheet hand-off — the Excel colleague's round trip, governed end to end.

send()     seals a snapshot of the approved assumptions (`spreadsheet_review_out`), builds an XLSX
           pack (Read me / Assumptions / Experience / Indications) and files it in the volume (out/).
receive()  takes the workbook back — an uploaded file today; Row Zero's `rz_` export table once Row
           Zero is connected to the workspace — files it as evidence (in/), lands the rows exactly as
           returned (`spreadsheet_review_in`) and checks them (`spreadsheet_review_changes`).
check      compares every row with the sealed snapshot: locked columns untouched, value parses and is
           in a plausible range, a rationale for every change, and no conflicting change to the
           approved value since the pack was sent.
apply()    turns accepted changes into one governed scenario per segment, calculated by the engine
           (the spreadsheet never produces the official number) and submitted for size-routed approval.

Every step is an audit event. Nothing in the spreadsheet becomes official without passing the checks.
"""
from __future__ import annotations
import asyncio
import csv
import hashlib
import io
import json
import re
import uuid
from datetime import datetime, timezone
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

import store
from config import CATALOG, SCHEMA, current_user, fqn
from indication_engine import ASSUMPTION_META, ASSUMPTION_ORDER
from sql import _client, execute_query
from store import _f, _i

DEFAULT_PERSONA = "Klaus Becker — Senior Pricing Actuary, Germany"
VOLUME = f"/Volumes/{CATALOG}/{SCHEMA}/spreadsheet_files"
ROWZERO_TABLE = "rz_assumption_review_return"     # Row Zero's export requires the rz_ prefix
SHEET = "Assumptions"
COLUMNS = ["handoff_id", "row_id", "product", "territory", "assumption", "unit",
           "current_value", "proposed_value", "rationale"]
EDITABLE = ("proposed_value", "rationale")

# Plausible range per assumption — a typo guard, not a pricing opinion (outside = rejected, range shown).
BOUNDS: dict[str, tuple[float, float]] = {
    "frequency_trend": (-0.15, 0.15), "severity_trend": (-0.10, 0.25),
    "loss_development_factor": (0.80, 3.00), "large_loss_load": (0.0, 0.50), "cat_load": (0.0, 0.50),
    "credibility": (0.0, 1.0), "experience_period_years": (3, 10),
    "expense_ratio": (0.0, 0.60), "commission_ratio": (0.0, 0.50),
    "reinsurance_load": (0.0, 0.30), "profit_provision": (-0.05, 0.25),
}
# Changes at or above these sizes are flagged as material in the check (still accepted).
MATERIAL = {"frequency_trend": 0.02, "severity_trend": 0.02, "loss_development_factor": 0.10}
PROVISIONS = ("expense_ratio", "commission_ratio", "reinsurance_load", "profit_provision")
MAX_PROVISIONS = 0.90   # provisions this high leave (almost) nothing for losses


class HandoffError(Exception):
    """A hand-off action that cannot proceed; the message is safe to show to the user."""


# ------------------------------------------------------------------ helpers
def _fmt(v: float | None, unit: str) -> str:
    if v is None:
        return "—"
    if unit == "pct":
        return f"{v:.2%}"
    if unit == "factor":
        return f"{v:.3f}"
    return f"{v:g}"


def _parse(raw: Any, unit: str) -> float:
    """Read a spreadsheet value the way a person typed it: 0.06, 6%, 6,0 % (German decimal comma)."""
    s = str(raw).strip().replace(" ", "").replace(" ", "")
    if "," in s and "." not in s:
        s = s.replace(",", ".")
    s = s.replace(",", "")
    is_pct = s.endswith("%")
    try:
        v = float(s[:-1]) / 100 if is_pct else float(s)
    except ValueError:
        raise ValueError(f"'{raw}' is not a number") from None
    if unit == "pct" and not is_pct and abs(v) >= 1:
        raise ValueError(f"'{raw}' looks like a percentage typed as a whole number — enter {v / 100:.4f} or {v:g}%")
    return v


def _cell(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float):
        return repr(v)
    return str(v).strip()


def _same(raw: str, ref: float | None) -> bool:
    if ref is None or raw == "":
        return False
    try:
        return abs(_parse(raw, "factor") - ref) < 1e-6      # unit-free compare: locked value as sent
    except ValueError:
        return False


def _seal(out_rows: list[dict]) -> str:
    canon = json.dumps(sorted([r["row_id"], round(r["current_value"], 10)] for r in out_rows))
    return hashlib.sha256(canon.encode()).hexdigest()[:16]


async def _insert_rows(table: str, cols: list[tuple[str, str]], rows: list[dict],
                       extra: dict[str, str] | None = None) -> None:
    """Insert many rows in ONE statement. Values travel as a single JSON parameter, never as SQL text;
    `extra` adds fixed SQL expressions (e.g. current_timestamp()) for audit columns."""
    if not rows:
        return
    struct = ", ".join(f"{c}: {t}" for c, t in cols)
    names = [c for c, _ in cols] + list((extra or {}).keys())
    exprs = [f"r.{c}" for c, _ in cols] + list((extra or {}).values())
    await execute_query(
        f"INSERT INTO {fqn(table)} ({', '.join(names)}) SELECT {', '.join(exprs)} "
        f"FROM (SELECT explode(from_json(:rows, 'array<struct<{struct}>>')) AS r)",
        {"rows": json.dumps(rows, default=str)})


def _upload(path: str, data: bytes) -> None:
    _client().files.upload(path, io.BytesIO(data), overwrite=True)


def _download(path: str) -> bytes:
    return _client().files.download(path).contents.read()


async def _handoff(hid: str) -> dict[str, Any]:
    rows = await execute_query(f"SELECT * FROM {fqn('spreadsheet_handoff')} WHERE handoff_id=:h", {"h": hid})
    if not rows:
        raise HandoffError(f"hand-off {hid} not found")
    return rows[0]


# ------------------------------------------------------------------ send
async def send(persona: str | None = None, period: int = 2027) -> dict[str, Any]:
    persona = (persona or DEFAULT_PERSONA).strip()
    base = await execute_query(
        f"""SELECT s.scenario_id, s.lob_code, s.territory_code, l.lob_label, t.territory_label,
                   s.experience_version, a.assumption_name, a.assumption_value
            FROM {fqn('indication_scenarios')} s
            JOIN {fqn('indication_assumptions')} a ON a.scenario_id = s.scenario_id
            JOIN {fqn('line_of_business')} l ON l.lob_code = s.lob_code
            JOIN {fqn('territory')} t ON t.territory_code = s.territory_code
            WHERE s.is_baseline AND s.indication_period = :period
            ORDER BY l.display_order, t.display_order""", {"period": period})
    if not base:
        raise HandoffError("no approved baselines found for this period")
    order = {n: i for i, n in enumerate(ASSUMPTION_ORDER)}
    seg_order: dict[tuple[str, str], int] = {}
    out: list[dict] = []
    for r in base:
        if r["assumption_name"] not in ASSUMPTION_META:
            continue
        key = (r["lob_code"], r["territory_code"])
        seg_order.setdefault(key, len(seg_order))
        meta = ASSUMPTION_META[r["assumption_name"]]
        out.append({
            "row_id": f"{r['lob_code']}|{r['territory_code']}|{r['assumption_name']}",
            "lob_code": r["lob_code"], "territory_code": r["territory_code"],
            "product": r["lob_label"], "territory": r["territory_label"],
            "assumption_name": r["assumption_name"], "assumption_label": meta["label"], "unit": meta["unit"],
            "current_value": _f(r["assumption_value"]), "baseline_scenario_id": r["scenario_id"],
        })
    out.sort(key=lambda r: (seg_order[(r["lob_code"], r["territory_code"])], order[r["assumption_name"]]))
    exp_ver = base[0]["experience_version"]
    now = datetime.now(timezone.utc)
    hid = f"RZ-{now:%Y%m%d-%H%M}-{uuid.uuid4().hex[:4]}"
    seal = _seal(out)
    for r in out:
        r["handoff_id"] = hid

    experience = await execute_query(
        f"""SELECT l.lob_label, t.territory_label, e.accident_year, e.earned_premium, e.reported_incurred,
                   e.claim_count, e.exposure, e.ldf_to_ultimate, e.rate_level_index
            FROM {fqn('indication_experience')} e
            JOIN {fqn('line_of_business')} l ON l.lob_code = e.lob_code
            JOIN {fqn('territory')} t ON t.territory_code = e.territory_code
            WHERE e.experience_version = :v
            ORDER BY l.display_order, t.display_order, e.accident_year""", {"v": exp_ver})
    pf = await store.portfolio(period)
    data = build_pack(hid, persona, current_user(), now, period, exp_ver, seal, out, experience, pf)
    pack_path = f"{VOLUME}/out/{hid}.xlsx"
    await asyncio.to_thread(_upload, pack_path, data)

    await execute_query(
        f"""INSERT INTO {fqn('spreadsheet_handoff')}
            (handoff_id, author_persona, sent_by, sent_at, indication_period, experience_version,
             row_count, seal, pack_path, status)
            VALUES (:h, :p, :u, current_timestamp(), :per, :ev, :n, :seal, :path, 'SENT')""",
        {"h": hid, "p": persona, "u": current_user(), "per": period, "ev": exp_ver,
         "n": len(out), "seal": seal, "path": pack_path})
    await _insert_rows("spreadsheet_review_out", [
        ("handoff_id", "STRING"), ("row_id", "STRING"), ("lob_code", "STRING"), ("territory_code", "STRING"),
        ("product", "STRING"), ("territory", "STRING"), ("assumption_name", "STRING"),
        ("assumption_label", "STRING"), ("unit", "STRING"), ("current_value", "DOUBLE"),
        ("baseline_scenario_id", "STRING")], out)
    await store._audit(hid, "HANDOFF_SENT", None, "SENT",
                       note=f"assumption review pack sent to spreadsheet for {persona}",
                       details={"rows": len(out), "segments": len(seg_order), "seal": seal,
                                "experience_version": exp_ver, "pack_path": pack_path,
                                "pack_sha256": hashlib.sha256(data).hexdigest()})
    return {"handoff_id": hid, "rows": len(out), "segments": len(seg_order), "seal": seal,
            "experience_version": exp_ver}


def build_pack(hid: str, persona: str, sent_by: str, sent_at: datetime, period: int, exp_ver: str,
               seal: str, out: list[dict], experience: list[dict], pf: dict[str, Any]) -> bytes:
    """The workbook the spreadsheet colleague receives. Only two columns on one sheet come back."""
    wb = openpyxl.Workbook()
    bold = Font(bold=True)
    yellow = PatternFill("solid", fgColor="FFF2CC")
    grey = PatternFill("solid", fgColor="EDEDED")

    rm = wb.active
    rm.title = "Read me"
    lines: list[tuple[str, str]] = [
        ("Bricksurance SE — assumption review pack", ""),
        ("", ""),
        ("Hand-off", hid),
        ("Prepared for", persona),
        ("Prepared by", sent_by),
        ("Prepared at (UTC)", f"{sent_at:%Y-%m-%d %H:%M}"),
        ("Indication period", str(period)),
        ("Experience data version", exp_ver),
        ("Seal (proves which approved values were sent)", seal),
        ("", ""),
        ("How to use this pack", ""),
        ("1.", "Work however you like: add sheets, formulas, pivots and charts. They stay yours."),
        ("2.", "On the Assumptions sheet, change proposed_value where you disagree with the approved value, "
               "and say why in rationale. A change without a rationale is not accepted."),
        ("3.", "Leave the grey columns as they are. They prove which approved value you reviewed."),
        ("4.", "Return the workbook in the Rate Indications app: Spreadsheet review, then Return. Every row is "
               "checked; accepted changes are recalculated by the governed engine and routed for approval."),
    ]
    for a, b in lines:
        rm.append([a, b])
    rm["A1"].font = Font(bold=True, size=14)
    rm["A11"].font = bold
    rm.column_dimensions["A"].width = 44
    rm.column_dimensions["B"].width = 110
    for row in rm.iter_rows(min_row=12, max_row=15):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")

    ws = wb.create_sheet(SHEET)
    ws.append(COLUMNS)
    for i, name in enumerate(COLUMNS, start=1):
        c = ws.cell(row=1, column=i)
        c.font = bold
        c.fill = yellow if name in EDITABLE else grey
    fmt = {"pct": "0.00%", "factor": "0.000", "years": "0"}
    for r in out:
        ws.append([r["handoff_id"], r["row_id"], r["product"], r["territory"], r["assumption_label"], r["unit"],
                   r["current_value"], r["current_value"], ""])
        n = ws.max_row
        ws.cell(row=n, column=7).number_format = fmt.get(r["unit"], "General")
        ws.cell(row=n, column=8).number_format = fmt.get(r["unit"], "General")
        ws.cell(row=n, column=8).fill = yellow
        ws.cell(row=n, column=9).fill = yellow
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:I{ws.max_row}"
    for col, w in zip("ABCDEFGHI", (22, 44, 22, 14, 24, 8, 14, 15, 60)):
        ws.column_dimensions[col].width = w

    ex = wb.create_sheet("Experience")
    ex.append(["product", "territory", "accident_year", "earned_premium", "reported_incurred",
               "reported_loss_ratio", "claim_count", "exposure", "ldf_to_ultimate", "rate_level_index"])
    for c in ex[1]:
        c.font = bold
    for r in experience:
        ex.append([r["lob_label"], r["territory_label"], _i(r["accident_year"]), _f(r["earned_premium"]),
                   _f(r["reported_incurred"]), None, _i(r["claim_count"]), _f(r["exposure"]),
                   _f(r["ldf_to_ultimate"]), _f(r["rate_level_index"])])
        n = ex.max_row
        ex.cell(row=n, column=6).value = f"=E{n}/D{n}"          # a real formula, as in their own workbook
        ex.cell(row=n, column=6).number_format = "0.0%"
        ex.cell(row=n, column=4).number_format = "#,##0"
        ex.cell(row=n, column=5).number_format = "#,##0"
    ex.freeze_panes = "A2"

    ind = wb.create_sheet("Indications")
    ind.append(["product", "territory", "approved_indicated_rate_change", "projected_loss_ratio",
                "on_level_earned_premium"])
    for c in ind[1]:
        c.font = bold
    for s in pf.get("segments", []):
        ind.append([s["lob_label"], s["territory_label"], s["baseline_indicated"], s["projected_loss_ratio"],
                    s["on_level_earned_premium"]])
        n = ind.max_row
        ind.cell(row=n, column=3).number_format = "+0.0%;-0.0%"
        ind.cell(row=n, column=4).number_format = "0.0%"
        ind.cell(row=n, column=5).number_format = "#,##0"

    wb.active = 1          # open on the Assumptions sheet
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def pack(hid: str) -> tuple[str, bytes]:
    h = await _handoff(hid)
    return f"{hid}.xlsx", await asyncio.to_thread(_download, h["pack_path"])


# ------------------------------------------------------------------ return
def _header_index(table: list[list[Any]]) -> int | None:
    for i, row in enumerate(table[:15]):
        names = {str(c).strip().lower() for c in row if c is not None}
        if {"row_id", "proposed_value"} <= names:
            return i
    return None


def read_rows(filename: str, data: bytes) -> list[dict[str, str]]:
    """The hand-off rows from a returned workbook. Extra sheets and columns are the author's own and ignored."""
    name = filename.lower()
    if name.endswith(".csv"):
        table: list[list[Any]] | None = list(csv.reader(io.StringIO(data.decode("utf-8-sig", errors="replace"))))
    elif name.endswith((".xlsx", ".xlsm")):
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
        table = None
        for sh in sorted(wb.worksheets, key=lambda s: s.title != SHEET):    # prefer the Assumptions sheet
            t = [list(row) for row in sh.iter_rows(values_only=True)]
            if _header_index(t) is not None:
                table = t
                break
        wb.close()
        if table is None:
            raise HandoffError("no sheet in the workbook has the hand-off columns (row_id, proposed_value)")
    else:
        raise HandoffError("upload the workbook as .xlsx or .csv")
    h = _header_index(table)
    if h is None:
        raise HandoffError("the file has no hand-off columns (row_id, proposed_value)")
    header = [str(c).strip().lower() if c is not None else "" for c in table[h]]
    idx = {c: i for i, c in enumerate(header) if c}
    rows = []
    for row in table[h + 1:]:
        def get(c: str) -> str:
            i = idx.get(c)
            return _cell(row[i]) if i is not None and i < len(row) else ""
        if get("row_id"):
            rows.append({c: get(c) for c in COLUMNS})
    if not rows:
        raise HandoffError("the hand-off sheet has no rows")
    return rows


async def receive(hid: str, filename: str, data: bytes) -> dict[str, Any]:
    h = await _handoff(hid)
    if h["status"] in ("APPLIED", "APPLYING"):
        raise HandoffError(f"{hid} is already being applied or applied — send a fresh pack for a new review")
    rows = read_rows(filename, data)
    sha = hashlib.sha256(data).hexdigest()
    safe = re.sub(r"[^A-Za-z0-9._-]+", "_", filename)[-80:]
    path = f"{VOLUME}/in/{hid}-{datetime.now(timezone.utc):%Y%m%dT%H%M%S}-{safe}"
    await asyncio.to_thread(_upload, path, data)
    return await _land_and_check(h, rows, "file", path, sha, filename)


async def pull_rowzero(hid: str) -> dict[str, Any]:
    """Read Row Zero's 'Export to Databricks' table — the connected path, no file on a laptop."""
    h = await _handoff(hid)
    if h["status"] in ("APPLIED", "APPLYING"):
        raise HandoffError(f"{hid} is already being applied or applied — send a fresh pack for a new review")
    exists = await execute_query(
        f"""SELECT 1 FROM `{CATALOG}`.information_schema.tables
            WHERE table_schema = :s AND table_name = :t""", {"s": SCHEMA, "t": ROWZERO_TABLE})
    if not exists:
        raise HandoffError(f"No Row Zero export found yet. In Row Zero, select the Assumptions table, choose "
                           f"Export to Databricks, schema {SCHEMA}, table name {ROWZERO_TABLE}.")
    raw = await execute_query(f"SELECT * FROM {fqn(ROWZERO_TABLE)}")
    rows = []
    for r in raw:
        low = {str(k).strip().lower(): v for k, v in r.items()}
        if not low.get("row_id"):
            continue
        rows.append({c: _cell(low.get(c)) for c in COLUMNS})
    rows = [r for r in rows if r["handoff_id"] in ("", hid)]
    if not rows:
        raise HandoffError(f"the Row Zero export has no rows for {hid} — export the pack sent with this hand-off")
    sha = hashlib.sha256(json.dumps(rows, sort_keys=True).encode()).hexdigest()
    return await _land_and_check(h, rows, "rowzero_export", f"{CATALOG}.{SCHEMA}.{ROWZERO_TABLE}", sha, ROWZERO_TABLE)


async def _land_and_check(h: dict, rows: list[dict], source: str, path: str, sha: str,
                          label: str) -> dict[str, Any]:
    hid = h["handoff_id"]
    await _insert_rows("spreadsheet_review_in", [
        ("handoff_id", "STRING"), ("row_id", "STRING"), ("product", "STRING"), ("territory", "STRING"),
        ("assumption", "STRING"), ("unit", "STRING"), ("current_value_returned", "STRING"),
        ("proposed_value_raw", "STRING"), ("rationale", "STRING"), ("source", "STRING"),
        ("source_path", "STRING"), ("source_sha256", "STRING"), ("loaded_by", "STRING")],
        [{"handoff_id": hid, "row_id": r["row_id"], "product": r["product"], "territory": r["territory"],
          "assumption": r["assumption"], "unit": r["unit"], "current_value_returned": r["current_value"],
          "proposed_value_raw": r["proposed_value"], "rationale": r["rationale"], "source": source,
          "source_path": path, "source_sha256": sha, "loaded_by": current_user()} for r in rows],
        extra={"loaded_at": "current_timestamp()"})
    await store._audit(hid, "HANDOFF_RETURNED", h["status"], "RETURNED",
                       note=f"workbook returned ({source}): {label}",
                       details={"source": source, "path": path, "sha256": sha, "rows": len(rows)})
    checked = await check(hid, rows)
    counts = {k: sum(1 for c in checked if c["check_status"] == k) for k in ("ACCEPTED", "REJECTED", "UNCHANGED")}
    await execute_query(
        f"""UPDATE {fqn('spreadsheet_handoff')} SET status='RETURNED', returned_by=:u,
            returned_at=current_timestamp(), return_source=:src, return_path=:p, return_sha256=:sha,
            accepted=:a, rejected=:r, unchanged=:n WHERE handoff_id=:h""",
        {"u": current_user(), "src": source, "p": path, "sha": sha, "a": counts["ACCEPTED"],
         "r": counts["REJECTED"], "n": counts["UNCHANGED"], "h": hid})
    await store._audit(hid, "HANDOFF_CHECKED", None, None,
                       note=f"{counts['ACCEPTED']} accepted, {counts['REJECTED']} rejected, "
                            f"{counts['UNCHANGED']} unchanged",
                       details={"rejected": [{"row_id": c["row_id"], "reason": c["reason"]}
                                             for c in checked if c["check_status"] == "REJECTED"]})
    return {"handoff_id": hid, **{k.lower(): v for k, v in counts.items()}}


async def check(hid: str, rows: list[dict]) -> list[dict]:
    """Check every returned row against the sealed snapshot and record the outcome with its reason."""
    snap = {r["row_id"]: r for r in await execute_query(
        f"SELECT * FROM {fqn('spreadsheet_review_out')} WHERE handoff_id=:h", {"h": hid})}
    live = await execute_query(
        f"""SELECT s.lob_code, s.territory_code, a.assumption_name, a.assumption_value
            FROM {fqn('indication_scenarios')} s
            JOIN {fqn('indication_assumptions')} a ON a.scenario_id = s.scenario_id
            WHERE s.scenario_id IN (SELECT DISTINCT baseline_scenario_id FROM {fqn('spreadsheet_review_out')}
                                    WHERE handoff_id = :h)""", {"h": hid})
    live_val = {f"{r['lob_code']}|{r['territory_code']}|{r['assumption_name']}": _f(r["assumption_value"])
                for r in live}
    out: list[dict] = []
    seen: set[str] = set()

    def rec(rid: str, s: dict | None, status: str, reason: str, proposed: float | None = None,
            rationale: str = "") -> dict:
        d = {"handoff_id": hid, "row_id": rid,
             "lob_code": s["lob_code"] if s else None, "territory_code": s["territory_code"] if s else None,
             "assumption_name": s["assumption_name"] if s else None,
             "current_value": _f(s["current_value"]) if s else None,
             "proposed_value": proposed, "rationale": rationale, "check_status": status, "reason": reason,
             "checked_by": current_user()}
        out.append(d)
        return d

    for r in rows:
        rid, s = r["row_id"], snap.get(r["row_id"])
        rationale = r["rationale"].strip()
        if s is None:
            rec(rid, None, "REJECTED", "row_id is not part of this hand-off (unknown or edited row_id)")
            continue
        if rid in seen:
            rec(rid, s, "REJECTED", "duplicate row — only the first copy of each row is read")
            continue
        seen.add(rid)
        if r["handoff_id"] and r["handoff_id"] != hid:
            rec(rid, s, "REJECTED", f"row belongs to hand-off {r['handoff_id']}, not {hid}")
            continue
        unit, cur = s["unit"], _f(s["current_value"])
        edited = [] if _same(r["current_value"], cur) else ["current_value"]
        for col, ref in (("product", s["product"]), ("territory", s["territory"]),
                         ("assumption", s["assumption_label"]), ("unit", unit)):
            if r[col] and r[col] != ref:
                edited.append(col)
        if edited:
            rec(rid, s, "REJECTED", f"locked column edited ({', '.join(edited)}) — the grey columns prove "
                                    f"which approved value was reviewed", rationale=rationale)
            continue
        if r["proposed_value"] == "":
            rec(rid, s, "UNCHANGED", "left blank — approved value kept")
            continue
        try:
            v = _parse(r["proposed_value"], unit)
        except ValueError as e:
            rec(rid, s, "REJECTED", str(e), rationale=rationale)
            continue
        if abs(v - cur) < 1e-9:
            rec(rid, s, "UNCHANGED", "no change", proposed=v)
            continue
        name = s["assumption_name"]
        lo, hi = BOUNDS.get(name, (float("-inf"), float("inf")))
        if not lo <= v <= hi:
            rec(rid, s, "REJECTED", f"{_fmt(v, unit)} is outside the plausible range "
                                    f"{_fmt(lo, unit)} to {_fmt(hi, unit)}", proposed=v, rationale=rationale)
            continue
        if unit == "years" and v != int(v):
            rec(rid, s, "REJECTED", "whole years only", proposed=v, rationale=rationale)
            continue
        if not rationale:
            rec(rid, s, "REJECTED", "changed without a rationale — say why in the rationale column", proposed=v)
            continue
        now = live_val.get(rid)
        if now is not None and abs(now - cur) > 1e-9:
            rec(rid, s, "REJECTED", f"the approved value changed to {_fmt(now, unit)} after this pack was sent "
                                    f"— send a fresh pack", proposed=v, rationale=rationale)
            continue
        big = abs(v - cur) >= MATERIAL.get(name, float("inf"))
        rec(rid, s, "ACCEPTED", f"{_fmt(cur, unit)} → {_fmt(v, unit)}"
                                + (" · material change" if big else ""), proposed=v, rationale=rationale)

    # Segment-level sense check: the provisions together must leave room for losses.
    by_seg: dict[tuple[str, str], list[dict]] = {}
    for c in out:
        if c["check_status"] == "ACCEPTED" and c["assumption_name"] in PROVISIONS:
            by_seg.setdefault((c["lob_code"], c["territory_code"]), []).append(c)
    for (lob, terr), items in by_seg.items():
        vals = {n: live_val.get(f"{lob}|{terr}|{n}", 0.0) or 0.0 for n in PROVISIONS}
        vals.update({c["assumption_name"]: c["proposed_value"] for c in items})
        if sum(vals.values()) >= MAX_PROVISIONS:
            for c in items:
                c["check_status"] = "REJECTED"
                c["reason"] = (f"expense, commission, reinsurance and profit together would be "
                               f"{sum(vals.values()):.0%} of premium — no room left for losses")

    await execute_query(f"DELETE FROM {fqn('spreadsheet_review_changes')} WHERE handoff_id=:h AND applied_scenario_id IS NULL",
                        {"h": hid})
    await _insert_rows("spreadsheet_review_changes", [
        ("handoff_id", "STRING"), ("row_id", "STRING"), ("lob_code", "STRING"), ("territory_code", "STRING"),
        ("assumption_name", "STRING"), ("current_value", "DOUBLE"), ("proposed_value", "DOUBLE"),
        ("rationale", "STRING"), ("check_status", "STRING"), ("reason", "STRING"), ("checked_by", "STRING")],
        out, extra={"checked_at": "current_timestamp()", "applied_scenario_id": "CAST(NULL AS STRING)"})
    return out


# ------------------------------------------------------------------ apply
# Applying runs in the background (one governed scenario per segment takes a few seconds each);
# the page polls detail() and watches the audit events arrive. Errors are kept here for the page.
_apply_errors: dict[str, str] = {}
_apply_tasks: dict[str, asyncio.Task] = {}


async def start_apply(hid: str) -> dict[str, Any]:
    h = await _handoff(hid)
    if h["status"] == "APPLYING" or (hid in _apply_tasks and not _apply_tasks[hid].done()):
        return {"handoff_id": hid, "status": "APPLYING"}
    if h["status"] != "RETURNED":
        raise HandoffError("return and check the workbook before applying" if h["status"] == "SENT"
                           else f"{hid} has already been applied")
    n = await execute_query(
        f"""SELECT count(*) n FROM {fqn('spreadsheet_review_changes')}
            WHERE handoff_id=:h AND check_status='ACCEPTED' AND applied_scenario_id IS NULL""", {"h": hid})
    if not n or not _i(n[0]["n"]):
        raise HandoffError("nothing was accepted, so there is nothing to apply")
    await execute_query(f"UPDATE {fqn('spreadsheet_handoff')} SET status='APPLYING' WHERE handoff_id=:h", {"h": hid})
    _apply_errors.pop(hid, None)
    _apply_tasks[hid] = asyncio.create_task(_apply_guarded(hid))     # inherits the request's user identity
    return {"handoff_id": hid, "status": "APPLYING"}


async def _apply_guarded(hid: str) -> None:
    try:
        await apply(hid)
    except Exception as e:  # noqa: BLE001
        import logging
        logging.getLogger("rate_indications").exception("apply %s failed", hid)
        _apply_errors[hid] = "Applying failed part-way; already-created scenarios are kept. See server logs."
        await execute_query(f"UPDATE {fqn('spreadsheet_handoff')} SET status='RETURNED' WHERE handoff_id=:h",
                            {"h": hid})
        await store._audit(hid, "HANDOFF_APPLY_FAILED", "APPLYING", "RETURNED", note=str(e)[:300])


async def apply(hid: str) -> dict[str, Any]:
    h = await _handoff(hid)
    if h["status"] not in ("RETURNED", "APPLYING"):
        raise HandoffError("return and check the workbook before applying" if h["status"] == "SENT"
                           else f"{hid} has already been applied")
    acc = await execute_query(
        f"""SELECT * FROM {fqn('spreadsheet_review_changes')}
            WHERE handoff_id=:h AND check_status='ACCEPTED' AND applied_scenario_id IS NULL""", {"h": hid})
    if not acc:
        raise HandoffError("nothing was accepted, so there is nothing to apply")
    period = _i(h["indication_period"])
    persona = h["author_persona"] or DEFAULT_PERSONA
    groups: dict[tuple[str, str], list[dict]] = {}
    for c in acc:
        groups.setdefault((c["lob_code"], c["territory_code"]), []).append(c)
    created = []
    for (lob, terr), items in groups.items():
        sid = await store.create_scenario(
            f"Spreadsheet review {hid} — {persona.split(' — ')[0]}", lob, terr, period, None,
            overrides={c["assumption_name"]: _f(c["proposed_value"]) for c in items})
        changes = [{"assumption": c["assumption_name"], "from": _f(c["current_value"]),
                    "to": _f(c["proposed_value"]), "rationale": c["rationale"]} for c in items]
        comment = "; ".join(f"{ASSUMPTION_META[c['assumption']]['label']}: {c['rationale']}" for c in changes)
        await execute_query(f"UPDATE {fqn('indication_scenarios')} SET comments=:c WHERE scenario_id=:sid",
                            {"c": f"From spreadsheet hand-off {hid} ({persona}). {comment}", "sid": sid})
        await store._audit(sid, "FROM_SPREADSHEET", None, None,
                           note=f"assumptions from spreadsheet hand-off {hid} (author: {persona})",
                           details={"handoff_id": hid, "changes": changes, "source": h["return_source"],
                                    "source_sha256": h["return_sha256"]})
        calc = await store.record_calculation(sid)
        await store.set_status(sid, "submit", note=f"submitted from spreadsheet hand-off {hid}")
        await execute_query(
            f"""UPDATE {fqn('spreadsheet_review_changes')} SET applied_scenario_id=:sid
                WHERE handoff_id=:h AND check_status='ACCEPTED'
                  AND array_contains(from_json(:ids, 'array<string>'), row_id)""",
            {"sid": sid, "h": hid, "ids": json.dumps([c["row_id"] for c in items])})
        new_ind = calc["result"]["indicated_rate_change"]
        created.append({"scenario_id": sid, "lob_code": lob, "territory_code": terr,
                        "indicated_rate_change": new_ind, "changes": len(items),
                        "requires": await store.required_role(abs(new_ind))})
    await execute_query(
        f"""UPDATE {fqn('spreadsheet_handoff')} SET status='APPLIED', applied_by=:u,
            applied_at=current_timestamp(), applied_scenarios=:s WHERE handoff_id=:h""",
        {"u": current_user(), "s": json.dumps([c["scenario_id"] for c in created]), "h": hid})
    await store._audit(hid, "HANDOFF_APPLIED", "RETURNED", "APPLIED",
                       note=f"{len(acc)} accepted change(s) applied as {len(created)} governed scenario(s), "
                            f"calculated and submitted for approval",
                       details={"scenarios": created})
    return {"handoff_id": hid, "scenarios": created}


# ------------------------------------------------------------------ reads
async def list_handoffs() -> list[dict]:
    rows = await execute_query(
        f"""SELECT handoff_id, author_persona, sent_by, sent_at, status, accepted, rejected, unchanged,
                   returned_at, applied_at
            FROM {fqn('spreadsheet_handoff')} ORDER BY sent_at DESC LIMIT 20""")
    for r in rows:
        for k in ("accepted", "rejected", "unchanged"):
            r[k] = _i(r[k])
    return rows


async def detail(hid: str) -> dict[str, Any]:
    h = await _handoff(hid)
    for k in ("accepted", "rejected", "unchanged", "row_count", "indication_period"):
        h[k] = _i(h[k])
    changes = await execute_query(
        f"""SELECT c.row_id, c.lob_code, c.territory_code, c.assumption_name, c.current_value, c.proposed_value,
                   c.rationale, c.check_status, c.reason, c.applied_scenario_id, o.product, o.territory,
                   o.assumption_label, o.unit
            FROM {fqn('spreadsheet_review_changes')} c
            LEFT JOIN {fqn('spreadsheet_review_out')} o ON o.handoff_id = c.handoff_id AND o.row_id = c.row_id
            WHERE c.handoff_id = :h
            ORDER BY CASE c.check_status WHEN 'ACCEPTED' THEN 0 WHEN 'REJECTED' THEN 1 ELSE 2 END, c.row_id""",
        {"h": hid})
    for c in changes:
        c["current_value"], c["proposed_value"] = _f(c["current_value"]), _f(c["proposed_value"])
    scen_ids = json.loads(h["applied_scenarios"]) if h.get("applied_scenarios") else []
    scenarios = []
    current = {x["scenario_id"]: x for x in await store.scenarios()} if scen_ids else {}
    for sid in scen_ids:
        s = current.get(sid)
        if not s:
            continue      # removed by a demo reset
        base = await store.scenario_detail(await store.baseline_scenario_id(s["lob_code"], s["territory_code"],
                                                                            s["indication_period"]))
        s["baseline_indicated"] = (base.get("result") or {}).get("indicated_rate_change")
        s["requires"] = await store.required_role(abs(s["indicated_rate_change"] or 0.0))
        scenarios.append(s)
    ids = [hid] + scen_ids
    events = await execute_query(
        f"""SELECT log_ts, scenario_id, action, actor, from_status, to_status, note
            FROM {fqn('indication_audit_log')}
            WHERE array_contains(from_json(:ids, 'array<string>'), scenario_id)
            ORDER BY log_ts""", {"ids": json.dumps(ids)})
    return {"handoff": h, "changes": changes, "scenarios": scenarios, "events": events,
            "rowzero_table": f"{SCHEMA}.{ROWZERO_TABLE}", "apply_error": _apply_errors.get(hid)}


async def reset() -> None:
    """Demo reset: clear hand-offs (the audit log keeps the record; files stay in the volume as evidence)."""
    for t in ("spreadsheet_review_changes", "spreadsheet_review_in", "spreadsheet_review_out", "spreadsheet_handoff"):
        await execute_query(f"DELETE FROM {fqn(t)}")
